"""Generate SYNTHETIC financial statements for the three fictional DealLens demo companies.

Each company is produced by a small three-statement model (income statement → balance sheet → cash
flow) so the statements are internally consistent: the balance sheet balances and the cash
roll-forward reconciles. Deliberate inconsistencies are injected only where listed in
DELIBERATE_ISSUES (and documented in samples/README.md) to demonstrate the validation engine.

Usage: processing/.venv/bin/python scripts/generate_synthetic_data.py [output_dir]
"""
from __future__ import annotations

import csv
import json
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

YEARS = [2022, 2023, 2024]
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "samples"


@dataclass
class Profile:
    name: str
    revenue: list[float]
    cogs_pct: list[float]
    sga_pct: float
    rnd_pct: float
    dso: list[float]
    dio: list[float]
    dpo: list[float]
    capex_pct: float
    da_pct: float
    debt: list[float]          # long-term debt balances
    st_debt: list[float]
    rate: float
    dividends_pct: float
    opening: dict


PROFILES = {
    "northstar": Profile("Northstar Components Ltd. (Synthetic)", [48000, 51400, 55100], [0.62, 0.62, 0.615], 0.17, 0.03,
                         [52, 53, 52], [68, 70, 69], [48, 49, 48], 0.05, 0.04, [14000, 13000, 12000], [1500, 1500, 1500], 0.06, 0.3,
                         {"cash": 4200, "ar": 6700, "inv": 5100, "oca": 900, "ppe": 21000, "onca": 3500, "ap": 3900, "accr": 1800, "ocl": 900, "oncl": 1200,
                          "ltd": 15000, "std": 1500}),
    "meridian": Profile("Meridian Consumer Brands Ltd. (Synthetic)", [82000, 94500, 109800], [0.62, 0.655, 0.69], 0.21, 0.0,
                        [45, 58, 72], [60, 76, 95], [40, 41, 38], 0.035, 0.03, [32000, 38000, 46000], [4000, 6000, 9000], 0.075, 0.1,
                        {"cash": 6500, "ar": 9800, "inv": 8000, "oca": 1500, "ppe": 26000, "onca": 9000, "ap": 5400, "accr": 2600, "ocl": 1400, "oncl": 2000,
                         "ltd": 30000, "std": 4000}),
    "apex": Profile("Apex Digital Services Ltd. (Synthetic)", [26000, 31200, 36500], [0.29, 0.28, 0.275], 0.31, 0.12,
                    [62, 60, 61], [0, 0, 0], [35, 34, 36], 0.02, 0.025, [0, 0, 0], [0, 0, 0], 0.0, 0.0,
                    {"cash": 9500, "ar": 4200, "inv": 0, "oca": 700, "ppe": 2100, "onca": 5200, "ap": 650, "accr": 1300, "ocl": 4100, "oncl": 300,
                     "ltd": 0, "std": 0}),
}

DELIBERATE_ISSUES = {
    "northstar": ["Income statement presents expenses as negative numbers in parentheses (sign-convention handling).",
                  "A 'Gross margin %' ratio row is present and is intentionally ignored (ratios are recalculated)."],
    "meridian": ["FY2024 'Total assets' is overstated by 150 versus its components (balance-sheet reconciliation variance).",
                 "Income statement contains an unmapped label 'Brand royalty pool' (unmapped line item).",
                 "Profile designed to trigger margin-deterioration, receivables/inventory-outpacing-revenue, OCF-below-EBITDA and working-capital red flags."],
    "apex": ["Reported EBITDA excludes share-based compensation (800/950/1,100), so it differs from EBIT + D&A (reported vs recalculated discrepancy).",
             "Statements are provided as a single CSV with stacked statements and a customer revenue schedule (customer concentration > 20%).",
             "No inventory (service business): DIO is unavailable and trade working capital treats inventory as nil."],
}


def model(p: Profile, key: str) -> dict:
    o = p.opening
    rows = {y: {} for y in YEARS}
    prev = dict(o)
    prev["equity"] = (o["cash"] + o["ar"] + o["inv"] + o["oca"] + o["ppe"] + o["onca"]) - (o["ap"] + o["accr"] + o["std"] + o["ocl"] + o["ltd"] + o["oncl"])
    for i, y in enumerate(YEARS):
        r = rows[y]
        rev = p.revenue[i]
        cogs = round(rev * p.cogs_pct[i])
        sga = round(rev * p.sga_pct)
        rnd = round(rev * p.rnd_pct)
        da = round(rev * p.da_pct)
        ebit = rev - cogs - sga - rnd - da
        interest = round((prev["ltd"] + prev["std"]) * p.rate)
        pbt = ebit - interest
        tax = round(max(pbt, 0) * 0.25)
        ni = pbt - tax
        r.update(revenue=rev, cogs=cogs, gp=rev - cogs, sga=sga, rnd=rnd, da=da, ebit=ebit, ebitda=ebit + da, interest=interest, pbt=pbt, tax=tax, ni=ni)
        ar = round(rev * p.dso[i] / 365)
        inv = round(cogs * p.dio[i] / 365)
        ap = round(cogs * p.dpo[i] / 365)
        oca = round(prev["oca"] * 1.04)
        accr = round(prev["accr"] * 1.05)
        ocl = round(prev["ocl"] * (1.15 if key == "apex" else 1.03))  # deferred revenue grows for Apex
        oncl, onca = prev["oncl"], prev["onca"]
        capex = round(rev * p.capex_pct)
        ppe = prev["ppe"] + capex - da
        ltd, std = p.debt[i], p.st_debt[i]
        div = round(max(ni, 0) * p.dividends_pct)
        ocf = ni + da - (ar - prev["ar"]) - (inv - prev["inv"]) - (oca - prev["oca"]) + (ap - prev["ap"]) + (accr - prev["accr"]) + (ocl - prev["ocl"])
        icf = -capex
        fcf_fin = (ltd - prev["ltd"]) + (std - prev["std"]) - div
        change = ocf + icf + fcf_fin
        cash = prev["cash"] + change
        equity = prev["equity"] + ni - div
        r.update(ar=ar, inv=inv, ap=ap, oca=oca, accr=accr, ocl=ocl, oncl=oncl, onca=onca, capex=capex, ppe=ppe, ltd=ltd, std=std, div=div,
                 ocf=ocf, icf=icf, fin=fcf_fin, change=change, cash=cash, begin_cash=prev["cash"], equity=equity,
                 tca=cash + ar + inv + oca, tcl=ap + accr + std + ocl)
        r["ta"] = r["tca"] + ppe + onca
        r["tl"] = r["tcl"] + ltd + oncl
        assert abs(r["ta"] - r["tl"] - equity) < 1e-6, (key, y)
        prev = {"cash": cash, "ar": ar, "inv": inv, "oca": oca, "ppe": ppe, "onca": onca, "ap": ap, "accr": accr, "ocl": ocl, "oncl": oncl,
                "ltd": ltd, "std": std, "equity": equity}
    return rows


def statements(key: str, m: dict) -> dict[str, list[tuple[str, list]]]:
    neg = key == "northstar"
    e = (lambda v: -v) if neg else (lambda v: v)
    col = lambda f: [f(m[y]) for y in YEARS]  # noqa: E731
    income = [("Revenue", col(lambda r: r["revenue"])), ("Cost of sales", col(lambda r: e(r["cogs"]))), ("Gross profit", col(lambda r: r["gp"]))]
    if key == "northstar":
        income.append(("Gross margin %", [f"{m[y]['gp'] / m[y]['revenue']:.1%}" for y in YEARS]))
    income += [("Selling, general and administrative expenses", col(lambda r: e(r["sga"])))]
    if key != "meridian":
        income.append(("Research and development", col(lambda r: e(r["rnd"]))))
    if key == "meridian":
        income.append(("Brand royalty pool", [0, 0, 0]))
    if key == "apex":
        income.append(("EBITDA", [m[y]["ebitda"] - sbc for y, sbc in zip(YEARS, (800, 950, 1100))]))
    income += [("Depreciation and amortisation", col(lambda r: e(r["da"]))), ("Operating profit", col(lambda r: r["ebit"])),
               ("Finance costs", col(lambda r: e(r["interest"]))), ("Profit before tax", col(lambda r: r["pbt"])),
               ("Income tax expense", col(lambda r: e(r["tax"]))), ("Net income", col(lambda r: r["ni"]))]
    ta = col(lambda r: r["ta"])
    if key == "meridian":
        ta[2] += 150
    balance = [("Cash and cash equivalents", col(lambda r: r["cash"])), ("Trade receivables", col(lambda r: r["ar"]))]
    if key != "apex":
        balance.append(("Inventories", col(lambda r: r["inv"])))
    balance += [("Prepayments", col(lambda r: r["oca"])), ("Total current assets", col(lambda r: r["tca"])),
                ("Property, plant and equipment", col(lambda r: r["ppe"])), ("Intangible assets", col(lambda r: r["onca"])),
                ("Total assets", ta), ("Trade payables", col(lambda r: r["ap"])), ("Accrued expenses", col(lambda r: r["accr"]))]
    if key != "apex":
        balance.append(("Short-term borrowings", col(lambda r: r["std"])))
    balance += [("Deferred revenue" if key == "apex" else "Other current liabilities", col(lambda r: r["ocl"])),
                ("Total current liabilities", col(lambda r: r["tcl"]))]
    if key != "apex":
        balance.append(("Long-term debt", col(lambda r: r["ltd"])))
    balance += [("Provisions", col(lambda r: r["oncl"])), ("Total liabilities", col(lambda r: r["tl"])), ("Total equity", col(lambda r: r["equity"])),
                ("Total liabilities and equity", col(lambda r: r["tl"] + r["equity"]))]
    cash_flow = [("Net income", col(lambda r: r["ni"])), ("Depreciation and amortisation", col(lambda r: r["da"])),
                 ("Net cash from operating activities", col(lambda r: r["ocf"])), ("Purchase of property, plant and equipment", col(lambda r: -r["capex"])),
                 ("Net cash used in investing activities", col(lambda r: r["icf"])), ("Net cash from financing activities", col(lambda r: r["fin"])),
                 ("Net increase in cash", col(lambda r: r["change"])), ("Cash at beginning of year", col(lambda r: r["begin_cash"])),
                 ("Cash at end of year", col(lambda r: r["cash"]))]
    return {"Income Statement": income, "Balance Sheet": balance, "Cash Flow Statement": cash_flow}


def write_xlsx(path: Path, company: str, sts: dict) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sts.items():
        ws = wb.create_sheet(title)
        ws.append([f"{company} — {title}"])
        ws.append(["SYNTHETIC DATA — fictional company for demonstration only"])
        ws.append(["USD in thousands"])
        ws.append([])
        ws.append(["Line item", *[f"FY{y}" for y in YEARS]])
        for c in ws[5]:
            c.font = Font(bold=True)
        for label, vals in rows:
            ws.append([label, *vals])
        ws.column_dimensions["A"].width = 48
    wb.save(path)


def write_csv(path: Path, company: str, sts: dict) -> None:
    with path.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow([f"{company} — Financial statements (SYNTHETIC)"])
        w.writerow(["USD '000"])
        w.writerow(["Line item", *[f"FY{y}" for y in YEARS]])
        for title, rows in sts.items():
            w.writerow([title])
            for label, vals in rows:
                w.writerow([label, *[f"{v:,.0f}" if isinstance(v, (int, float)) else v for v in vals]])


def write_text_pdf(path: Path, company: str, title: str, rows) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    _, h = A4
    y = h - 60
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, f"{company} - {title}")
    y -= 16
    c.setFont("Helvetica", 9)
    c.drawString(50, y, "SYNTHETIC DATA - USD in thousands")
    y -= 24
    c.setFont("Helvetica-Bold", 9)
    for i, yr in enumerate(YEARS):
        c.drawRightString(380 + i * 70, y, f"FY{yr}")
    c.setFont("Helvetica", 9)
    for label, vals in rows:
        if any(isinstance(v, str) for v in vals):
            continue
        y -= 15
        c.drawString(50, y, label)
        for i, v in enumerate(vals):
            c.drawRightString(380 + i * 70, y, f"({abs(v):,.0f})" if v < 0 else f"{v:,.0f}")
    c.showPage()
    c.save()


def write_scanned_pdf(path: Path) -> None:
    """An image-only PDF (no text layer), used to verify the unsupported-scan path."""
    c = canvas.Canvas(str(path), pagesize=A4)
    for i in range(30):
        c.rect(50, 780 - i * 20, 200 + (i * 37) % 250, 8, fill=1, stroke=0)
    c.showPage()
    c.save()


def write_customer_schedule(path: Path, revenue: list[float]) -> None:
    shares = [0.28, 0.14, 0.09, 0.07, 0.05]
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue by customer"
    ws.append(["Apex Digital Services Ltd. (Synthetic) — Revenue by customer"])
    ws.append(["USD in thousands — SYNTHETIC"])
    ws.append(["Customer", *[f"FY{y}" for y in YEARS]])
    for i, s in enumerate(shares):
        ws.append([f"Customer {chr(65 + i)} (fictional)", *[round(r * s) for r in revenue]])
    ws.append(["All other customers", *[round(r * (1 - sum(shares))) for r in revenue]])
    ws.append(["Total revenue", *[round(r) for r in revenue]])
    wb.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    manifest = {}
    for key, profile in PROFILES.items():
        m = model(profile, key)
        sts = statements(key, m)
        files = []
        if key == "apex":
            write_csv(OUT / "apex-financial-statements.csv", profile.name, sts)
            write_customer_schedule(OUT / "apex-revenue-by-customer.xlsx", profile.revenue)
            files += ["apex-financial-statements.csv", "apex-revenue-by-customer.xlsx"]
        else:
            write_xlsx(OUT / f"{key}-financial-statements.xlsx", profile.name, sts)
            files.append(f"{key}-financial-statements.xlsx")
        if key == "northstar":
            write_text_pdf(OUT / "northstar-income-statement.pdf", profile.name, "Income Statement", sts["Income Statement"])
            files.append("northstar-income-statement.pdf")
        manifest[key] = {"company": profile.name, "files": files, "deliberate_issues": DELIBERATE_ISSUES[key],
                         "model": {str(y): m[y] for y in YEARS}}
    write_scanned_pdf(OUT / "scanned-statement-unsupported.pdf")
    (OUT / "manifest.json").write_text(json.dumps(manifest, indent=2))
    print(f"Wrote synthetic samples to {OUT}")


if __name__ == "__main__":
    main()
