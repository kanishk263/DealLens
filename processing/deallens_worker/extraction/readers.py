"""Format readers that turn a verified document into cell grids, with resource limits.

Spreadsheets are read with cached values only; formulas are never evaluated and macros are never run
(macro-enabled workbooks are rejected earlier by signature checks)."""
from __future__ import annotations

import csv
import io
import warnings

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..errors import MalformedDocument, UnsupportedDocument
from .models import Cell, Grid

MAX_SHEETS = 30
MAX_ROWS = 5000
MAX_COLS = 60
MAX_PDF_PAGES = 60
MIN_PDF_TEXT_CHARS = 40


def read_xlsx(data: bytes) -> list[Grid]:
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            values_wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            formula_wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:  # openpyxl raises many exception types for corrupt files
        raise MalformedDocument("CORRUPT_WORKBOOK", "The workbook could not be read; it may be corrupt or password-protected.") from exc
    if len(values_wb.sheetnames) > MAX_SHEETS:
        raise MalformedDocument("TOO_MANY_SHEETS", f"Workbooks may contain at most {MAX_SHEETS} worksheets.")
    grids = []
    for name in values_wb.sheetnames:
        ws, wf = values_wb[name], formula_wb[name]
        if getattr(ws, "sheet_state", "visible") != "visible":
            continue  # hidden sheets are skipped deliberately and reported in metadata
        rows: list[list[Cell]] = []
        for r_idx, (vrow, frow) in enumerate(zip(ws.iter_rows(max_col=MAX_COLS), wf.iter_rows(max_col=MAX_COLS)), start=1):
            if r_idx > MAX_ROWS:
                raise MalformedDocument("SHEET_TOO_LARGE", f"Worksheet '{name}' exceeds {MAX_ROWS} rows.")
            row = []
            for c_idx, (vc, fc) in enumerate(zip(vrow, frow), start=1):
                value = getattr(vc, "value", None)
                formula = getattr(fc, "value", None)
                missing = isinstance(formula, str) and formula.startswith("=") and value is None
                if value is None and not missing:
                    continue
                row.append(Cell(value=value, ref=f"{get_column_letter(c_idx)}{r_idx}", row=r_idx, col=c_idx, formula_missing_value=missing))
            rows.append(row)
        grids.append(Grid(name=name, rows=rows, sheet=name))
    values_wb.close()
    formula_wb.close()
    return grids


def hidden_sheets(data: bytes) -> list[str]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(io.BytesIO(data), read_only=True)
    names = [n for n in wb.sheetnames if getattr(wb[n], "sheet_state", "visible") != "visible"]
    wb.close()
    return names


def read_xls(data: bytes) -> list[Grid]:
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception as exc:
        raise MalformedDocument("CORRUPT_WORKBOOK", "The .xls workbook could not be read.") from exc
    if book.nsheets > MAX_SHEETS:
        raise MalformedDocument("TOO_MANY_SHEETS", f"Workbooks may contain at most {MAX_SHEETS} worksheets.")
    grids = []
    for sheet in book.sheets():
        if sheet.visibility != 0:
            continue
        if sheet.nrows > MAX_ROWS:
            raise MalformedDocument("SHEET_TOO_LARGE", f"Worksheet '{sheet.name}' exceeds {MAX_ROWS} rows.")
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(min(sheet.ncols, MAX_COLS)):
                cell = sheet.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    value = "#VALUE!"
                row.append(Cell(value=value, ref=f"{get_column_letter(c + 1)}{r + 1}", row=r + 1, col=c + 1))
            rows.append(row)
        grids.append(Grid(name=sheet.name, rows=rows, sheet=sheet.name))
    book.release_resources()
    return grids


def read_csv(data: bytes, name: str) -> list[Grid]:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise MalformedDocument("CSV_ENCODING", "The CSV file is not UTF-8 or Windows-1252 encoded.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for r_idx, raw in enumerate(csv.reader(io.StringIO(text), dialect), start=1):
        if r_idx > MAX_ROWS:
            raise MalformedDocument("CSV_TOO_LARGE", f"CSV files may contain at most {MAX_ROWS} rows.")
        rows.append([Cell(value=v, ref=f"R{r_idx}C{c}", row=r_idx, col=c) for c, v in enumerate(raw[:MAX_COLS], start=1) if v.strip() != ""])
    return [Grid(name=name, rows=rows)]


def read_pdf(data: bytes) -> list[Grid]:
    """Text-based PDFs only. Image-only (scanned) PDFs are rejected explicitly; OCR is out of scope."""
    import pdfplumber
    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as exc:
        raise MalformedDocument("CORRUPT_PDF", "The PDF could not be opened; it may be corrupt or encrypted.") from exc
    with pdf:
        if len(pdf.pages) > MAX_PDF_PAGES:
            raise MalformedDocument("PDF_TOO_LARGE", f"PDFs may contain at most {MAX_PDF_PAGES} pages.")
        grids: list[Grid] = []
        text_chars = 0
        for p_no, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            text_chars += len(text.strip())
            tables = page.extract_tables()
            if tables:
                for t_no, table in enumerate(tables, start=1):
                    rows = [[Cell(value=v, ref=f"p{p_no}:t{t_no}:R{r}C{c}", row=r, col=c) for c, v in enumerate(row, start=1) if v not in (None, "")]
                            for r, row in enumerate(table, start=1)]
                    grids.append(Grid(name=f"Page {p_no} table {t_no}", rows=rows, page=p_no, table_ref=f"page {p_no}, table {t_no}"))
            elif text.strip():
                grids.append(Grid(name=f"Page {p_no} text", rows=_text_rows(text, p_no), page=p_no, table_ref=f"page {p_no}, text lines"))
        if text_chars < MIN_PDF_TEXT_CHARS:
            raise UnsupportedDocument(
                "SCANNED_PDF_UNSUPPORTED",
                "This PDF contains no extractable text (it appears to be scanned). OCR is not supported in this release; "
                "upload the original spreadsheet or a text-based PDF export.",
            )
        return grids


def _text_rows(text: str, page: int) -> list[list[Cell]]:
    """Split text lines into a label followed by trailing numeric tokens (e.g. 'Revenue 1,200 1,050')."""
    import re

    from .periods import parse_period
    token = re.compile(r"\(?-?[$£€]?\d[\d,]*(?:\.\d+)?\)?-?|—|–")
    rows = []
    for r, line in enumerate(text.splitlines(), start=1):
        parts = line.strip().split()
        values: list[str] = []
        while parts and (token.fullmatch(parts[-1]) or parse_period(parts[-1]) is not None):
            values.insert(0, parts.pop())
        label = " ".join(parts)
        cells = []
        if label:
            cells.append(Cell(value=label, ref=f"p{page}:line{r}:C1", row=r, col=1))
        for c, v in enumerate(values, start=2):
            cells.append(Cell(value=v, ref=f"p{page}:line{r}:C{c}", row=r, col=c))
        if not label and values:
            # Header lines consisting only of years, right-aligned under value columns.
            cells = [Cell(value=v, ref=f"p{page}:line{r}:C{c}", row=r, col=c) for c, v in enumerate(values, start=2)]
        rows.append(cells)
    return _align_numeric_columns(rows)


def _align_numeric_columns(rows: list[list[Cell]]) -> list[list[Cell]]:
    """Right-align value tokens so the last token of every line shares the same column index."""
    width = max((len(r) for r in rows), default=0)
    aligned = []
    for r in rows:
        if not r:
            aligned.append(r)
            continue
        label = [c for c in r if c.col == 1]
        values = [c for c in r if c.col != 1]
        shift = width - 1 - len(values)
        aligned.append(label + [Cell(value=c.value, ref=c.ref, row=c.row, col=2 + shift + i) for i, c in enumerate(values)])
    return aligned
